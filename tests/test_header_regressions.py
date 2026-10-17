#!/usr/bin/env python3
"""Synthetic DOC header/footer regressions, including destructive-output checks.

Uses the selected headless LibreOffice runtime, never an installed Office app.
Inputs and reports are written below --work-dir; no real user document is used.
"""
from __future__ import annotations

import argparse
import copy
import hashlib
import json
import os
from pathlib import Path
import re
import struct
import sys
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from lxml import etree

HERE = Path(__file__).resolve().parent
REPO = HERE.parent
sys.path.insert(0, str(REPO / 'engine'))
import converter

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
PKG = 'http://schemas.openxmlformats.org/package/2006/relationships'
N = {'w': W, 'r': R}
DEFAULT_ENGINE = Path.home() / '.cache/codex-runtimes/codex-primary-runtime/dependencies/native/libreoffice-headless/libreoffice/LibreOfficeDev.app/Contents/MacOS/soffice'
PAGE_FIELD = re.compile(r'^\s*(PAGE|NUMPAGES)(?=\s|\\|$)', re.I)


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def indexed_header_checks():
    """Exercise PlcfHdd boundaries without using an Office-generated fixture."""
    literal_header = '页眉正文\x03保留控制字符及辅助平面字形😀'
    authored_separator = '\x03自定义脚注分隔线'
    slots = [authored_separator, '\x04', '', '\x03', '\x04', '',
             '', literal_header, '', '静态页脚 88', '', '']
    text = ''
    cps = [0]
    for story in slots:
        text += story + '\r' if story else ''
        cps.append(len(text.encode('utf-16le')) // 2)
    text += '\r'  # Header Document's final guard paragraph, outside all slots.
    cps.append(0)  # MS-DOC declares the final CP undefined.
    plc = struct.pack('<%dI' % len(cps), *cps)
    sed = bytearray(20)
    struct.pack_into('<II', sed, 0, 0, 1)
    struct.pack_into('<I', sed, 10, 0xffffffff)
    table = plc + bytes(sed) + b'\0'
    word = bytearray(260)
    for index, offset, size in [(11, 0, len(plc)), (6, len(plc), len(sed)),
                                 (31, len(plc) + len(sed), 1)]:
        struct.pack_into('<II', word, index * 8, offset, size)
    parsed = converter.binary_header_parts(word, table, 0, text)
    checks = {'literal_header_control_retained': parsed['sections'][0]['header']['default'] == literal_header,
              'separator_control_scoped_to_separator': parsed['separators'][0] == authored_separator,
              'supplementary_character_retained': '😀' in parsed['header'][1]}
    corrupted = bytearray(table)
    struct.pack_into('<I', corrupted, 4, 0xffffffff)
    try:
        converter.binary_header_parts(word, corrupted, 0, text)
    except converter.ConversionError:
        checks['corrupt_header_index_rejected'] = True
    else:
        checks['corrupt_header_index_rejected'] = False
    return checks


def append_field(paragraph, instruction, cached):
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), instruction)
    run = OxmlElement('w:r')
    text = OxmlElement('w:t')
    text.text = cached
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


def write_story(story, text):
    story.is_linked_to_previous = False
    story.paragraphs[0].text = text


def page(doc, marker):
    doc.add_paragraph(marker)
    doc.add_paragraph('正文普通空格  保留，不换行\u00a0空格保留；编号 2026-17，数值 88。')


def make_fixture(name):
    doc = Document()
    doc.styles['Normal'].font.name = 'Times New Roman'
    doc.styles['Normal'].font.size = Pt(12)
    doc.styles['Normal']._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Songti SC')
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.1)
    section.left_margin = section.right_margin = Cm(2.2)
    write_story(section.header, '页眉甲|普通  空格|不换行\u00a0空格')
    write_story(section.footer, '页脚乙|静态数字 88')
    page(doc, f'{name} 正文第一页')
    doc.add_page_break()
    page(doc, f'{name} 正文第二页')

    if name in ('first-even-odd', 'same-text-variants'):
        section.different_first_page_header_footer = True
        doc.settings.odd_and_even_pages_header_footer = True
        for variant, label in [('header', '奇数页眉甲'), ('footer', '奇数页脚乙'),
                               ('even_page_header', '偶数页眉丙'), ('even_page_footer', '偶数页脚丁'),
                               ('first_page_header', '首页页眉戊'), ('first_page_footer', '首页页脚己')]:
            if name == 'same-text-variants':
                label = '所有变体页眉' if variant.endswith('header') else '所有变体页脚'
            write_story(getattr(section, variant), label + '|不得丢失')
        doc.add_page_break()
        page(doc, f'{name} 正文第三页')
    elif name in ('sections-distinct', 'sections-shared', 'sections-mixed', 'sections-clear'):
        section2 = doc.add_section(WD_SECTION_START.NEW_PAGE)
        if name != 'sections-shared':
            write_story(section2.footer, '第二节页脚丁|静态数字 99')
        if name == 'sections-distinct':
            write_story(section2.header, '第二节页眉丙|不可挪到第一节')
        if name == 'sections-clear':
            write_story(section2.header, '')
            write_story(section2.footer, '')
        page(doc, f'{name} 第二节正文第一页')
        doc.add_page_break()
        page(doc, f'{name} 第二节正文第二页')
    elif name == 'repeated-text':
        write_story(section.header, '重复页眉甲|重复页眉甲|结束')
        section.header.add_paragraph('整段重复乙')
        section.header.add_paragraph('整段重复乙')
        write_story(section.footer, '重复页脚丙|重复页脚丙|静态 88')
    elif name == 'page-fields':
        write_story(section.header, '档案编号 2026-17|静态页码 9')
        write_story(section.footer, '静态页码 88|第 ')
        p = section.footer.paragraphs[0]
        append_field(p, 'PAGE', '1')
        p.add_run(' 页，共 ')
        append_field(p, 'NUMPAGES', '2')
        p.add_run(' 页|固定编号 99')
    elif name == 'separator-only':
        section.header.paragraphs[0].text = ''
        section.footer.paragraphs[0].text = ''
    return doc


def add_note_parts(path, actual_notes):
    """Create standard separators; unused separators also occur in real DOC."""
    with ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    document = etree.fromstring(parts['word/document.xml'])
    relationships = etree.fromstring(parts['word/_rels/document.xml.rels'])
    content_types = etree.fromstring(parts['[Content_Types].xml'])
    for kind, label in [('footnote', '脚注'), ('endnote', '尾注')]:
        root = etree.Element(qn('w:' + kind + 's'), nsmap={'w': W})
        for note_id, typ, marker in [('-1', 'separator', 'separator'),
                                     ('0', 'continuationSeparator', 'continuationSeparator')]:
            note = etree.SubElement(root, qn('w:' + kind), {qn('w:id'): note_id, qn('w:type'): typ})
            paragraph = etree.SubElement(note, qn('w:p'))
            run = etree.SubElement(paragraph, qn('w:r'))
            etree.SubElement(run, qn('w:' + marker))
        if actual_notes:
            note = etree.SubElement(root, qn('w:' + kind), {qn('w:id'): '2'})
            p = etree.SubElement(note, qn('w:p'))
            etree.SubElement(etree.SubElement(p, qn('w:r')), qn('w:t')).text = f'{label}正文乙|不能丢失'
            run = etree.SubElement(document.find('.//w:body/w:p', N), qn('w:r'))
            etree.SubElement(run, qn('w:' + kind + 'Reference'), {qn('w:id'): '2'})
        name = kind + 's.xml'
        parts['word/' + name] = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
        etree.SubElement(relationships, '{%s}Relationship' % PKG,
                         {'Id': 'rIdRegression' + kind, 'Type': R + '/' + kind + 's', 'Target': name})
        etree.SubElement(content_types, '{http://schemas.openxmlformats.org/package/2006/content-types}Override',
                         {'PartName': '/word/' + name, 'ContentType': 'application/vnd.openxmlformats-officedocument.wordprocessingml.' + kind + 's+xml'})
    parts['word/document.xml'] = etree.tostring(document, xml_declaration=True, encoding='UTF-8', standalone=True)
    parts['word/_rels/document.xml.rels'] = etree.tostring(relationships, xml_declaration=True, encoding='UTF-8', standalone=True)
    parts['[Content_Types].xml'] = etree.tostring(content_types, xml_declaration=True, encoding='UTF-8', standalone=True)
    with ZipFile(path, 'w', ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


def story_text(root):
    """Independent static-text reader; only PAGE/NUMPAGES caches are excluded."""
    chunks = []
    fields = []
    for node in root.iter():
        if node.tag == qn('w:fldChar'):
            typ = node.get(qn('w:fldCharType'))
            if typ == 'begin':
                fields.append({'instruction': '', 'result': False})
            elif typ == 'separate':
                assert fields, 'field separator without begin'
                fields[-1]['result'] = True
            elif typ == 'end':
                assert fields, 'field end without begin'
                fields.pop()
        elif node.tag == qn('w:instrText'):
            if fields:
                fields[-1]['instruction'] += node.text or ''
        elif node.tag in (qn('w:t'), qn('w:tab')):
            if any(not item['result'] or PAGE_FIELD.match(item['instruction']) for item in fields):
                continue
            if any(ancestor.tag == qn('w:fldSimple') and PAGE_FIELD.match(ancestor.get(qn('w:instr'), ''))
                   for ancestor in node.iterancestors()):
                continue
            chunks.append('\t' if node.tag == qn('w:tab') else node.text or '')
    assert not fields, 'unclosed field'
    return ''.join(chunks)


def document_snapshot(path):
    """Resolve effective section stories, including links to previous sections."""
    with ZipFile(path) as archive:
        document = etree.fromstring(archive.read('word/document.xml'))
        settings = etree.fromstring(archive.read('word/settings.xml'))
        rels = etree.fromstring(archive.read('word/_rels/document.xml.rels'))
        targets = {node.get('Id'): node.get('Target') for node in rels.findall('{%s}Relationship' % PKG)}
        even_and_odd = settings.find('w:evenAndOddHeaders', N) is not None
        inherited = {}
        sections = []
        for properties in document.findall('.//w:sectPr', N):
            active = ['default'] + (['even'] if even_and_odd else [])
            if properties.find('w:titlePg', N) is not None:
                active.append('first')
            for kind in ('header', 'footer'):
                for reference in properties.findall('w:' + kind + 'Reference', N):
                    target = targets[reference.get('{%s}id' % R)]
                    part = target.lstrip('/') if target.startswith('/') else 'word/' + target
                    inherited[kind, reference.get(qn('w:type'))] = etree.fromstring(archive.read(part))
            stories = {}
            for kind in ('header', 'footer'):
                for variant in active:
                    root = inherited.get((kind, variant))
                    stories[f'{kind}:{variant}'] = story_text(root) if root is not None else ''
            sections.append(stories)
        sizes = [{key: node.get(qn('w:' + key)) for key in ('w', 'h')}
                 for node in document.findall('.//w:sectPr/w:pgSz', N)]
        instructions = []
        for name in archive.namelist():
            if re.fullmatch(r'word/(?:header|footer)\d*\.xml', name):
                root = etree.fromstring(archive.read(name))
                instructions.extend((node.text or '').strip().split()[0].upper()
                                    for node in root.findall('.//w:instrText', N) if (node.text or '').strip())
                instructions.extend(node.get(qn('w:instr'), '').strip().split()[0].upper()
                                    for node in root.findall('.//w:fldSimple', N) if node.get(qn('w:instr'), '').strip())
        notes = {}
        for kind in ('footnote', 'endnote'):
            part = 'word/' + kind + 's.xml'
            if part not in archive.namelist():
                continue
            root = etree.fromstring(archive.read(part))
            notes[kind] = [story_text(note) for note in root.findall('w:' + kind, N)
                           if note.get(qn('w:type')) not in ('separator', 'continuationSeparator', 'continuationNotice')]
            # Empty built-in separators need not survive as explicit parts;
            # authored separator text is content and must remain intact.
            notes[kind + '_authored_separators'] = sorted(story_text(note) for note in root.findall('w:' + kind, N)
                if note.get(qn('w:type')) in ('separator', 'continuationSeparator', 'continuationNotice') and story_text(note))
        notes = {key: value for key, value in notes.items() if value}
        return {'body': story_text(document.find('w:body', N)), 'sections': sections,
                'page_sizes': sizes, 'notes': notes,
                'dynamic_fields': sorted(set(i for i in instructions if i in ('PAGE', 'NUMPAGES')))}


def mutate_package(source, target, mutation):
    with ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    roots = {name: etree.fromstring(data) for name, data in parts.items()
             if re.fullmatch(r'word/(?:header|footer)\d*\.xml', name)}
    mutation(roots, parts)
    for name, root in roots.items():
        parts[name] = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
    with ZipFile(target, 'w', ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


def replace_literal(roots, old, new):
    for root in roots.values():
        nodes = root.findall('.//w:t', N)
        combined = ''.join(node.text or '' for node in nodes)
        start = combined.find(old)
        if start < 0:
            continue
        end = start + len(old)
        position = 0
        inserted = False
        for node in nodes:
            text = node.text or ''
            stop = position + len(text)
            if stop > start and position < end:
                prefix = text[:max(0, start - position)]
                suffix = text[max(0, end - position):]
                node.text = prefix + (new if not inserted else '') + suffix
                inserted = True
            position = stop
        return
    raise AssertionError(f'mutation target not found: {old}')


def replace_part_xml(parts, name, mutate):
    root = etree.fromstring(parts[name])
    mutate(root)
    parts[name] = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)


def swap_variant_references(_roots, parts):
    def swap(root):
        references = root.findall('.//w:sectPr/w:headerReference', N)
        odd = next(node for node in references if node.get(qn('w:type')) == 'default')
        even = next(node for node in references if node.get(qn('w:type')) == 'even')
        odd_id, even_id = odd.get('{%s}id' % R), even.get('{%s}id' % R)
        odd.set('{%s}id' % R, even_id)
        even.set('{%s}id' % R, odd_id)
    replace_part_xml(parts, 'word/document.xml', swap)


def swap_section_references(_roots, parts):
    def swap(root):
        sections = root.findall('.//w:sectPr', N)
        first = sections[0].find("w:headerReference[@w:type='default']", N)
        second = sections[1].find("w:headerReference[@w:type='default']", N)
        first_id, second_id = first.get('{%s}id' % R), second.get('{%s}id' % R)
        first.set('{%s}id' % R, second_id)
        second.set('{%s}id' % R, first_id)
    replace_part_xml(parts, 'word/document.xml', swap)


def disable_even_headers(_roots, parts):
    def disable(root):
        node = root.find('w:evenAndOddHeaders', N)
        assert node is not None
        node.set(qn('w:val'), '0')
    replace_part_xml(parts, 'word/settings.xml', disable)


def delete_body_character(_roots, parts):
    def delete(root):
        node = root.find('.//w:body/w:p/w:r/w:t', N)
        assert node is not None and node.text
        node.text = node.text[:-1]
    replace_part_xml(parts, 'word/document.xml', delete)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--engine', type=Path, default=Path(os.environ.get('LO_ENGINE', str(DEFAULT_ENGINE))))
    parser.add_argument('--work-dir', type=Path, default=HERE / '_work' / 'header-regressions')
    args = parser.parse_args()
    engine = args.engine.expanduser().resolve()
    if not engine.is_file():
        parser.error('LibreOffice engine not found.')
    work = args.work_dir.expanduser().resolve()
    work.mkdir(parents=True, exist_ok=True)
    report = {'parser': indexed_header_checks(), 'positive': {}, 'negative': {}, 'all_inputs_unchanged': True}
    artifacts = {}
    names = ['distinct', 'first-even-odd', 'same-text-variants', 'sections-distinct', 'sections-shared', 'sections-clear',
             'sections-mixed', 'repeated-text', 'page-fields', 'separator-only', 'notes-and-separators']
    for name in names:
        folder = work / name
        folder.mkdir(exist_ok=True)
        fixture = folder / 'fixture.docx'
        make_fixture(name).save(fixture)
        if name in ('separator-only', 'notes-and-separators'):
            add_note_parts(fixture, actual_notes=name == 'notes-and-separators')
        doc = converter.convert_engine(engine, fixture, folder / 'binary', folder / 'export-profile', 'doc:MS Word 97', 90)
        before = sha(doc)
        result = converter.convert_one(doc, folder / 'converted', engine)
        record = {'status': result['status'], 'message': result['message'], 'source_unchanged': before == sha(doc)}
        report['all_inputs_unchanged'] &= record['source_unchanged']
        if result['output']:
            output = Path(result['output'])
            expected = document_snapshot(fixture)
            actual = document_snapshot(output)
            differences = {key: {'expected': expected[key], 'actual': actual[key]}
                           for key in expected if expected[key] != actual[key]}
            record['independent_content_and_section_checks'] = not differences
            record['differences'] = differences
            artifacts[name] = (doc, output)
        else:
            record['independent_content_and_section_checks'] = False
        report['positive'][name] = record
        print(f"{name}: {record['status']}; independent={record['independent_content_and_section_checks']}", flush=True)

    # Mutate real DOCX engine output, then run the full publish gate with exactly
    # that malformed candidate. ODT still comes from the original binary DOC.
    # This proves that fixing false failures does not merely bypass auditing.
    literal_cases = [
        ('lost-header', 'distinct', '页眉甲', ''),
        ('changed-footer', 'distinct', '页脚乙', '篡改页脚'),
        ('lost-repeat', 'repeated-text', '重复页眉甲|重复页眉甲', '重复页眉甲'),
        ('changed-static-page-number', 'page-fields', '静态页码 88', '静态页码 89'),
        ('changed-identifier', 'page-fields', '2026-17', '2026-18'),
    ]
    negative_cases = [(label, name, lambda roots, _parts, old=old, new=new: replace_literal(roots, old, new))
                      for label, name, old, new in literal_cases]
    negative_cases.extend([
        ('swapped-even-odd-header-references', 'first-even-odd', swap_variant_references),
        ('swapped-section-header-references', 'sections-distinct', swap_section_references),
        ('disabled-even-header-setting-with-identical-text', 'same-text-variants', disable_even_headers),
        ('lost-body-character', 'distinct', delete_body_character),
    ])
    for label, fixture_name, mutation in negative_cases:
        if fixture_name not in artifacts:
            report['negative'][label] = {'passed': False, 'reason': 'positive fixture failed'}
            continue
        doc, output = artifacts[fixture_name]
        folder = work / label
        folder.mkdir(exist_ok=True)
        malformed = folder / 'tampered.docx'
        mutate_package(output, malformed, mutation)
        baseline = converter.convert_engine(engine, doc, folder / 'baseline', folder / 'baseline-profile', 'odt:writer8', 90)
        destination = folder / 'rejected-output'
        def engine_candidate(_engine, _source, _output, _profile, spec, _timeout):
            return baseline if spec.startswith('odt:') else malformed
        source_hash = sha(doc)
        with patch.object(converter, 'convert_engine', side_effect=engine_candidate):
            result = converter.convert_one(doc, destination, engine)
        passed = result['status'] == 'error' and result['output'] is None
        passed &= not destination.exists() or not any(destination.iterdir())
        passed &= sha(doc) == source_hash
        report['negative'][label] = {'passed': passed, 'message': result['message']}
        print(f'{label}: correctly rejected={passed}', flush=True)

    if 'distinct' in artifacts:
        from content_audit import audit_extra
        doc, output = artifacts['distinct']
        folder = work / 'header-control-not-ignored'
        baseline = converter.convert_engine(engine, doc, folder / 'baseline', folder / 'profile', 'odt:writer8', 90)
        kind, stories = converter.detect(doc.read_bytes())
        stories = copy.deepcopy(stories)
        metadata = stories['header_parts']
        original = metadata['sections'][0]['header']['default']
        metadata['sections'][0]['header']['default'] = original + '\x03'
        metadata['header'] = [text + '\x03' if text == original else text for text in metadata['header']]
        result = audit_extra(kind, stories, baseline, output)
        report['negative']['header-control-not-ignored'] = {'passed': not result['passed'], 'message': result['message']}
        print(f"header-control-not-ignored: correctly rejected={not result['passed']}", flush=True)
        # LO's DOC fixture exporter itself discards custom separator text. Use
        # independent binary-story metadata to test this import-loss boundary,
        # rather than claiming that an empty exported separator proves fidelity.
        _, stories = converter.detect(doc.read_bytes())
        stories = copy.deepcopy(stories)
        stories['header_parts']['separators'][0] += '\x03自定义脚注分隔线文字'
        result = audit_extra(kind, stories, baseline, output)
        report['negative']['authored-separator-loss-rejected'] = {'passed': not result['passed'], 'message': result['message']}
        print(f"authored-separator-loss-rejected: correctly rejected={not result['passed']}", flush=True)

    passed = (report['all_inputs_unchanged'] and all(report['parser'].values()) and
              all(item['status'] in ('ok', 'warning') and item['independent_content_and_section_checks']
                  for item in report['positive'].values()) and
              all(item['passed'] for item in report['negative'].values()))
    report['passed'] = passed
    report_path = work / 'header-regression-results.json'
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(f'Report: {report_path}')
    raise SystemExit(0 if passed else 1)


if __name__ == '__main__':
    main()
