#!/usr/bin/env python3
"""Synthetic DOC hyphen roundtrips and fail-closed content mutations.

All text is fabricated. Word's U+001E/U+001F markers must retain their visible
hyphen semantics in body, header/footer, and note stories. Ordinary whitespace,
static numbers, punctuation, and character order remain exact content.
"""
from __future__ import annotations

import argparse
import copy
import hashlib
import json
import os
from pathlib import Path
import sys
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / 'engine'))
import converter
import content_audit
from test_header_regressions import add_note_parts

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
N = {'w': W}
SPECIAL = {'noBreakHyphen': '\u2011', 'softHyphen': '\u00ad'}
DEFAULT_ENGINE = Path.home() / '.cache/codex-runtimes/codex-primary-runtime/dependencies/native/libreoffice-headless/libreoffice/LibreOfficeDev.app/Contents/MacOS/soffice'


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def load_parts(path):
    with ZipFile(path) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def save_parts(parts, path):
    with ZipFile(path, 'w', ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


def xml_bytes(root):
    return etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)


def append_content(paragraph, label, representation):
    paragraph.add_run(f'{label}|编号 2026-17|ASCII A-B|普通空格  双格|不换行 A\u00a0B|全角 A\u3000B|不可断 E')
    for name, after in [('noBreakHyphen', 'F|可选 G'), ('softHyphen', 'H|顺序甲乙丙丁。')]:
        if representation == 'elements':
            paragraph.add_run()._r.append(OxmlElement('w:' + name))
        else:
            paragraph.add_run(SPECIAL[name])
        paragraph.add_run(after)


def make_fixture(path, representation):
    doc = Document()
    doc.styles['Normal'].font.name = 'Times New Roman'
    doc.styles['Normal'].font.size = Pt(12)
    doc.styles['Normal']._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Songti SC')
    append_content(doc.add_paragraph(), '正文合成样本', representation)
    append_content(doc.sections[0].header.paragraphs[0], '页眉合成样本', representation)
    append_content(doc.sections[0].footer.paragraphs[0], '页脚合成样本', representation)
    doc.save(path)
    add_note_parts(path, actual_notes=True)
    parts = load_parts(path)
    for kind, label in [('footnote', '脚注合成样本'), ('endnote', '尾注合成样本')]:
        root = etree.fromstring(parts[f'word/{kind}s.xml'])
        paragraph = root.find(f'w:{kind}[@w:id="2"]/w:p', N)
        paragraph.clear()
        # Build the exact same synthetic content using python-docx's paragraph
        # object, then move only its runs into the note paragraph.
        scratch = Document().add_paragraph()
        append_content(scratch, label, representation)
        for run in scratch._p:
            paragraph.append(copy.deepcopy(run))
        parts[f'word/{kind}s.xml'] = xml_bytes(root)
    save_parts(parts, path)


def static_text(root):
    """Independent literal/semantic text snapshot for these field-free fixtures."""
    chunks = []
    for node in root.iter():
        if node.tag == qn('w:t'):
            chunks.append(node.text or '')
        elif node.tag == qn('w:noBreakHyphen'):
            chunks.append('\u2011')
        elif node.tag == qn('w:softHyphen'):
            chunks.append('\u00ad')
        elif node.tag == qn('w:tab'):
            chunks.append('\t')
    return ''.join(chunks)


def snapshot(path):
    parts = load_parts(path)
    document = etree.fromstring(parts['word/document.xml'])
    result = {'body': static_text(document.find('w:body', N))}
    for kind in ('header', 'footer'):
        result[kind] = sorted(text for name, data in parts.items()
            if name.startswith('word/' + kind) and name.endswith('.xml')
            if (text := static_text(etree.fromstring(data))))
    for kind in ('footnote', 'endnote'):
        root = etree.fromstring(parts[f'word/{kind}s.xml'])
        result[kind] = [static_text(node) for node in root.findall('w:' + kind, N)
            if node.get(qn('w:type')) not in ('separator', 'continuationSeparator', 'continuationNotice')]
    return result


def candidate_representation(source, target, form):
    """Change equivalent OOXML spelling, preserving independently read content."""
    parts = load_parts(source)
    for name, data in list(parts.items()):
        if not name.startswith('word/') or not name.endswith('.xml'):
            continue
        root = etree.fromstring(data)
        if form == 'unicode':
            for marker, char in SPECIAL.items():
                for node in root.findall('.//w:' + marker, N):
                    node.tag = qn('w:t')
                    node.text = char
        elif form == 'elements':
            for node in list(root.findall('.//w:t', N)):
                value = node.text or ''
                if not any(char in value for char in SPECIAL.values()):
                    continue
                parent, index = node.getparent(), node.getparent().index(node)
                parent.remove(node)
                for char in value:
                    marker = next((key for key, special in SPECIAL.items() if special == char), None)
                    replacement = etree.Element(qn('w:' + marker)) if marker else etree.Element(qn('w:t'), dict(node.attrib))
                    if marker is None:
                        replacement.text = char
                        replacement.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
                    parent.insert(index, replacement)
                    index += 1
        parts[name] = xml_bytes(root)
    save_parts(parts, target)


def mutate(source, target, part_kind, marker=None, replacement=None, old=None, new=None):
    parts = load_parts(source)
    name = 'word/document.xml' if part_kind == 'body' else next(
        name for name in parts if name.startswith('word/' + part_kind) and name.endswith('.xml'))
    root = etree.fromstring(parts[name])
    if marker:
        node = root.find('.//w:' + marker, N)
        assert node is not None, f'{marker} absent in {name}'
        if replacement is None:
            node.getparent().remove(node)
        else:
            node.tag, node.text = qn('w:t'), replacement
    else:
        node = next(node for node in root.findall('.//w:t', N) if old in (node.text or ''))
        node.text = node.text.replace(old, new, 1)
    parts[name] = xml_bytes(root)
    save_parts(parts, target)


def reader_checks(fixture):
    text = '甲\x1e乙\x1f丙|A-B|  |\u00a0|\u3000|88'
    expected = '甲\u2011乙\u00ad丙|A-B|  |\u00a0|\u3000|88'
    checks = {'binary_body_semantics': converter.visible_binary(text) == expected,
              'binary_extra_semantics': content_audit._binary_visible(text) == expected,
              'xml_body_elements': converter.docx_info(fixture.read_bytes())['main'].replace('\r', '') == snapshot(fixture)['body'],
              'ascii_hyphen_distinct': converter.visible_binary('A\x1eB') != converter.visible_binary('A-B')}
    for label, root_tag in [('header', 'hdr'), ('footer', 'ftr'), ('footnote', 'footnote'), ('textbox', 'txbxContent')]:
        root = etree.fromstring(f'<w:{root_tag} xmlns:w="{W}"><w:p><w:r><w:t>甲</w:t><w:noBreakHyphen/><w:t>乙</w:t><w:softHyphen/><w:t>丙  \u00a0\u300088</w:t></w:r></w:p></w:{root_tag}>'.encode())
        checks[label + '_elements'] = content_audit._word_visible(root) == '甲\u2011乙\u00ad丙  \u00a0\u300088'
    return checks


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--engine', type=Path, default=Path(os.environ.get('LO_ENGINE', str(DEFAULT_ENGINE))))
    parser.add_argument('--work-dir', type=Path, default=HERE / '_work' / 'body-regressions')
    args = parser.parse_args()
    engine, work = args.engine.expanduser().resolve(), args.work_dir.expanduser().resolve()
    if not engine.is_file():
        parser.error('LibreOffice engine not found.')
    work.mkdir(parents=True, exist_ok=True)
    report = {'reader': {}, 'positive': {}, 'negative': {}, 'all_inputs_unchanged': True}
    artifacts = {}
    for form in ('elements', 'unicode'):
        folder = work / ('source-' + form)
        folder.mkdir(exist_ok=True)
        fixture = folder / 'synthetic.docx'
        make_fixture(fixture, form)
        if form == 'elements':
            report['reader'] = reader_checks(fixture)
        binary = converter.convert_engine(engine, fixture, folder / 'binary', folder / 'export-profile', 'doc:MS Word 97', 90)
        source_hash = sha(binary)
        stories = converter.read_binary_word(binary.read_bytes())
        has_markers = all('\x1e' in stories[kind] and '\x1f' in stories[kind]
                          for kind in ('main', 'header', 'footnote', 'endnote'))
        result = converter.convert_one(binary, folder / 'converted', engine)
        source_unchanged = sha(binary) == source_hash
        independent = bool(result['output']) and snapshot(fixture) == snapshot(Path(result['output']))
        report['all_inputs_unchanged'] &= source_unchanged
        report['positive']['source-' + form] = {'status': result['status'], 'message': result['message'],
            'binary_markers_present': has_markers, 'independent_story_checks': independent,
            'source_unchanged': source_unchanged}
        print(f'source-{form}: {result["status"]}; independent={independent}; DOC markers={has_markers}', flush=True)
        if result['output']:
            artifacts[form] = (binary, Path(result['output']))

    if 'elements' in artifacts:
        binary, output = artifacts['elements']
        baseline = converter.convert_engine(engine, binary, work / 'baseline', work / 'baseline-profile', 'odt:writer8', 90)
        candidates = {}
        for form in ('elements', 'unicode'):
            candidate = work / (form + '-candidate.docx')
            candidate_representation(output, candidate, form)
            candidates[form] = candidate
            def engine_candidate(_engine, _source, _output, _profile, spec, _timeout):
                return baseline if spec.startswith('odt:') else candidate
            with patch.object(converter, 'convert_engine', side_effect=engine_candidate):
                result = converter.convert_one(binary, work / ('equivalent-' + form), engine)
            independent = snapshot(output) == snapshot(candidate)
            report['positive']['candidate-' + form] = {'status': result['status'], 'message': result['message'],
                'independent_story_checks': independent}
            print(f'candidate-{form}: {result["status"]}; independent={independent}', flush=True)

        mutations = [
            ('lost-body-nobreak', 'body', {'marker': 'noBreakHyphen'}),
            ('body-nobreak-changed-to-ascii', 'body', {'marker': 'noBreakHyphen', 'replacement': '-'}),
            ('lost-body-soft', 'body', {'marker': 'softHyphen'}),
            ('body-soft-changed-to-ascii', 'body', {'marker': 'softHyphen', 'replacement': '-'}),
            ('lost-body-character', 'body', {'old': '丙', 'new': ''}),
            ('lost-ordinary-space', 'body', {'old': '  双格', 'new': ' 双格'}),
            ('nbsp-changed-to-space', 'body', {'old': '\u00a0', 'new': ' '}),
            ('ideographic-space-changed', 'body', {'old': '\u3000', 'new': ' '}),
            ('changed-static-number', 'body', {'old': '2026-17', 'new': '2026-18'}),
            ('changed-character-order', 'body', {'old': '甲乙丙丁', 'new': '甲丙乙丁'}),
        ]
        mutations.extend((f'lost-{kind}-{marker}', kind, {'marker': marker})
                         for kind in ('header', 'footer', 'footnote', 'endnote')
                         for marker in ('noBreakHyphen', 'softHyphen'))
        for label, part_kind, mutation in mutations:
            candidate = work / (label + '.docx')
            mutate(candidates['elements'], candidate, part_kind, **mutation)
            destination = work / ('rejected-' + label)
            def engine_candidate(_engine, _source, _output, _profile, spec, _timeout):
                return baseline if spec.startswith('odt:') else candidate
            source_hash = sha(binary)
            with patch.object(converter, 'convert_engine', side_effect=engine_candidate):
                result = converter.convert_one(binary, destination, engine)
            unchanged = sha(binary) == source_hash
            passed = result['status'] == 'error' and result['output'] is None and unchanged
            passed &= not destination.exists() or not any(destination.iterdir())
            report['all_inputs_unchanged'] &= unchanged
            report['negative'][label] = {'passed': passed, 'message': result['message']}
            print(f'{label}: correctly rejected={passed}', flush=True)

    passed = (all(report['reader'].values()) and report['all_inputs_unchanged'] and len(report['positive']) == 4 and
        len(report['negative']) == 18 and all(record['status'] in ('ok', 'warning') and record['independent_story_checks']
        and record.get('binary_markers_present', True) for record in report['positive'].values()) and
        all(record['passed'] for record in report['negative'].values()))
    report['passed'] = passed
    path = work / 'body-regression-results.json'
    path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(f'Report: {path}', flush=True)
    raise SystemExit(0 if passed else 1)


if __name__ == '__main__':
    main()
