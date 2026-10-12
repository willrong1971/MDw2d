from pathlib import Path
import argparse
import os
from xml.sax.saxutils import escape

parser = argparse.ArgumentParser(description='Generate synthetic document fixtures; no user documents are used.')
parser.add_argument('--output-dir', type=Path, default=Path(__file__).resolve().parent / '_work' / 'fixtures')
parser.add_argument('--engine', default=os.environ.get('LO_ENGINE', str(Path.home() / '.cache/codex-runtimes/codex-primary-runtime/dependencies/native/libreoffice-headless/libreoffice/LibreOfficeDev.app/Contents/MacOS/soffice')))
args = parser.parse_args()
ROOT = args.output_dir.expanduser().resolve()
SOFFICE = str(Path(args.engine).expanduser().resolve())
if not Path(SOFFICE).is_file():
    parser.error('LibreOffice engine not found; pass --engine or set LO_ENGINE.')
ROOT.mkdir(parents=True, exist_ok=True)
from docx import Document
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image, ImageDraw
import json, zipfile, shutil, hashlib, subprocess, os
from lxml import etree

NS_W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
NS_R = 'http://schemas.openxmlformats.org/package/2006/relationships'
NS_CT = 'http://schemas.openxmlformats.org/package/2006/content-types'

# The packaged headless runtime uses Fontconfig rather than native macOS font
# discovery. Scope its configuration/cache to these test intermediates.
(ROOT/'fontconfig-cache').mkdir(exist_ok=True)
(ROOT/'fonts.conf').write_text('''<?xml version="1.0"?>
<!DOCTYPE fontconfig SYSTEM "fonts.dtd">
<fontconfig>
 <dir>/System/Library/Fonts</dir>
 <dir>/System/Library/Fonts/Supplemental</dir>
 <dir>/Library/Fonts</dir>
 <dir>'''+escape(str(Path(SOFFICE).parent.parent/'Resources/fonts/truetype'))+'''</dir>
 <cachedir>'''+escape(str(ROOT/'fontconfig-cache'))+'''</cachedir>
</fontconfig>
''',encoding='utf-8')
lo_env=os.environ.copy(); lo_env['FONTCONFIG_FILE']=str(ROOT/'fonts.conf')

im = Image.new('RGB', (320, 100), 'white')
draw = ImageDraw.Draw(im)
for x, color in [(10, '#435e7a'), (115, '#849db0'), (220, '#cbd4dc')]:
    draw.rectangle((x, 10, x+90, 90), fill=color, outline='black', width=2)
draw.line((100, 50, 115, 50), fill='black', width=3)
draw.line((205, 50, 220, 50), fill='black', width=3)
im.save(ROOT/'fixture-image.png')

def base_doc():
    doc=Document()
    sec=doc.sections[0]
    sec.page_width=Cm(21); sec.page_height=Cm(29.7)
    sec.top_margin=Cm(2.1); sec.bottom_margin=Cm(2.1)
    sec.left_margin=Cm(2.2); sec.right_margin=Cm(2.2)
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
        style=doc.styles[name]
        style.font.name='Times New Roman'
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Songti SC')
        style.font.color.rgb=RGBColor(0,0,0)
        for border in style._element.xpath('./w:pPr/w:pBdr'):
            border.getparent().remove(border)
    doc.styles['Normal'].font.size=Pt(12)
    doc.styles['Normal'].paragraph_format.space_after=Pt(5)
    return doc

doc=base_doc()
doc.core_properties.title='通用文档转换测试基准'
doc.core_properties.author='Conversion fixture'
doc.add_paragraph('通用文档转换测试基准',style='Title')
doc.add_paragraph('此文档用于检查格式识别、文字保留和常见版式元素。全部文字、图形和数据均为合成测试内容。')
doc.add_paragraph('文字与空白',style='Heading 1')
doc.add_paragraph('中文甲乙丙，English ABC，数字 0123456789，标点：；？！“引号”（括号）。')
doc.add_paragraph('字间  两个空格|不换行\u00a0空格|全角\u3000空格|首尾空白  ')
p=doc.add_paragraph('混合格式：')
p.add_run('粗体').bold=True
p.add_run('、')
p.add_run('斜体').italic=True
p.add_run('、')
p.add_run('下划线').underline=True
p.add_run('、普通文字。')
p=doc.add_paragraph('显式换行前')
p.add_run().add_break()
p.add_run('显式换行后\t制表符后')
p=doc.add_paragraph('脚注测试位置')
r=OxmlElement('w:r'); rPr=OxmlElement('w:rPr'); rStyle=OxmlElement('w:rStyle'); rStyle.set(qn('w:val'),'FootnoteReference'); rPr.append(rStyle); r.append(rPr)
fnref=OxmlElement('w:footnoteReference'); fnref.set(qn('w:id'),'2'); r.append(fnref); p._p.append(r)
p.add_run('，正文继续。')
p=doc.add_paragraph('批注测试文字')
comment_start=OxmlElement('w:commentRangeStart'); comment_start.set(qn('w:id'),'0')
comment_end=OxmlElement('w:commentRangeEnd'); comment_end.set(qn('w:id'),'0')
p._p.insert(0, comment_start); p._p.append(comment_end)
r=OxmlElement('w:r'); cref=OxmlElement('w:commentReference'); cref.set(qn('w:id'),'0'); r.append(cref); p._p.append(r)
doc.add_paragraph('表格与图片',style='Heading 1')
t=doc.add_table(rows=3,cols=3); t.style='Table Grid'
values=[['项目','数量','说明'],['中文项目甲','12','保留单元格顺序'],['中文项目乙','3.50','不可修改数字']]
for row,vrow in zip(t.rows,values):
    for cell,v in zip(row.cells,vrow): cell.text=v
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.add_run().add_picture(str(ROOT/'fixture-image.png'),width=Inches(3.2))
doc.add_paragraph('图 1 三块矩形示意图')
doc.add_page_break()
doc.add_paragraph('第二页与编号',style='Heading 1')
doc.add_paragraph('第一项说明。',style='List Number')
doc.add_paragraph('第二项说明。',style='List Number')
doc.add_paragraph('结束段落：文字内容不得增删改写，允许转换引擎造成小幅排版差异。')
sec=doc.sections[0]
sec.header.paragraphs[0].text='转换测试页眉'
f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER
f.add_run('转换测试页脚 第 ')
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE')
r=OxmlElement('w:r'); txt=OxmlElement('w:t'); txt.text='1'; r.append(txt); fld.append(r); f._p.append(fld)
f.add_run(' 页')
base_path=ROOT/'baseline.docx'; doc.save(base_path)

# Add true OOXML footnote and comment parts, with the required relationships.
with zipfile.ZipFile(base_path) as z: parts={i.filename:z.read(i.filename) for i in z.infolist()}
footnotes=etree.Element(qn('w:footnotes'),nsmap={'w':NS_W})
for fid,type_,child in [('-1','separator','separator'),('0','continuationSeparator','continuationSeparator')]:
    fn=etree.SubElement(footnotes,qn('w:footnote')); fn.set(qn('w:id'),fid); fn.set(qn('w:type'),type_)
    p=etree.SubElement(fn,qn('w:p')); r=etree.SubElement(p,qn('w:r')); etree.SubElement(r,qn('w:'+child))
fn=etree.SubElement(footnotes,qn('w:footnote')); fn.set(qn('w:id'),'2')
p=etree.SubElement(fn,qn('w:p')); r=etree.SubElement(p,qn('w:r')); txt=etree.SubElement(r,qn('w:t')); txt.text='脚注正文：这是一条独立脚注，编号不可进入正文。'
comments=etree.Element(qn('w:comments'),nsmap={'w':NS_W})
comment=etree.SubElement(comments,qn('w:comment')); comment.set(qn('w:id'),'0'); comment.set(qn('w:author'),'测试作者'); comment.set(qn('w:date'),'2026-10-12T00:00:00Z')
p=etree.SubElement(comment,qn('w:p')); r=etree.SubElement(p,qn('w:r')); txt=etree.SubElement(r,qn('w:t')); txt.text='批注正文：此批注用于测试结构保留。'
parts['word/footnotes.xml']=etree.tostring(footnotes,xml_declaration=True,encoding='UTF-8',standalone=True)
parts['word/comments.xml']=etree.tostring(comments,xml_declaration=True,encoding='UTF-8',standalone=True)
rels=etree.fromstring(parts['word/_rels/document.xml.rels'])
for rid,target,rt in [('rIdFixtureFootnotes','footnotes.xml',RT.FOOTNOTES),('rIdFixtureComments','comments.xml',RT.COMMENTS)]:
    rel=etree.SubElement(rels,'{%s}Relationship'%NS_R); rel.set('Id',rid); rel.set('Type',rt); rel.set('Target',target)
parts['word/_rels/document.xml.rels']=etree.tostring(rels,xml_declaration=True,encoding='UTF-8',standalone=True)
ct=etree.fromstring(parts['[Content_Types].xml'])
for part,ctype in [('/word/footnotes.xml','application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml'),('/word/comments.xml','application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml')]:
    ov=etree.SubElement(ct,'{%s}Override'%NS_CT); ov.set('PartName',part); ov.set('ContentType',ctype)
parts['[Content_Types].xml']=etree.tostring(ct,xml_declaration=True,encoding='UTF-8',standalone=True)
with zipfile.ZipFile(base_path,'w',zipfile.ZIP_DEFLATED) as z:
    for name,data in parts.items(): z.writestr(name,data)

plain=base_doc(); plain.add_paragraph('简单有效文档',style='Title'); plain.add_paragraph('正文内容应原样保留：张三  李四，2026年10月12日。'); plain.save(ROOT/'plain.docx')

exports=[]
for fmt in ['doc:MS Word 97','rtf:Rich Text Format','odt:writer8']:
    suffix=fmt.split(':')[0]
    profile=ROOT/('lo-profile-'+suffix)
    cmd=[SOFFICE,'-env:UserInstallation='+profile.as_uri(),'--headless','--convert-to',fmt,'--outdir',str(ROOT),str(base_path)]
    result=subprocess.run(cmd,capture_output=True,text=True,timeout=60,env=lo_env)
    exports.append({'format':fmt,'returncode':result.returncode,'stdout':result.stdout.strip(),'stderr':result.stderr.strip()})
    if not (ROOT/('baseline.'+suffix)).is_file(): raise RuntimeError(str(exports[-1]))
shutil.copyfile(ROOT/'baseline.doc',ROOT/'binary-disguised.wps')
shutil.copyfile(base_path,ROOT/'ooxml-disguised.wps')
shutil.copyfile(ROOT/'baseline.rtf',ROOT/'rtf-disguised.doc')
shutil.copyfile(base_path,ROOT/'ooxml-disguised.doc')
(ROOT/'broken.docx').write_bytes(base_path.read_bytes()[:200])
(ROOT/'unsupported.wps').write_text('这只是纯文本，不是 WPS 文档。不能因为后缀是 .wps 就当成文档转换。\n',encoding='utf-8')
with zipfile.ZipFile(ROOT/'invalid-package.docx','w',zipfile.ZIP_DEFLATED) as z: z.writestr('readme.txt','ZIP 格式但没有 Word 文档结构。')

w_ns={'w':NS_W}
xml_names=['word/document.xml','word/header1.xml','word/footer1.xml','word/footnotes.xml','word/comments.xml']
expected_parts={name:etree.fromstring(parts[name]).xpath('//w:t/text()',namespaces=w_ns) for name in xml_names}
(ROOT/'baseline-text.json').write_text(json.dumps(expected_parts,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
positive=[]
for name,detected,action in [
    ('baseline.docx','ooxml','copy'),('plain.docx','ooxml','copy'),
    ('ooxml-disguised.wps','ooxml','copy'),('ooxml-disguised.doc','ooxml','copy'),
    ('baseline.doc','ole-word','convert'),('binary-disguised.wps','ole-word','convert'),
    ('baseline.rtf','rtf','convert'),('rtf-disguised.doc','rtf','convert'),('baseline.odt','odt','convert')
]:
    positive.append({'filename':name,'detected_family':detected,'expected_success':True,'expected_action':action,'reference_docx':'plain.docx' if name=='plain.docx' else 'baseline.docx','byte_identical_expected':action=='copy'})
neg=[{'filename':'broken.docx','expected_success':False,'reason':'truncated ZIP must not report successful conversion'}, {'filename':'unsupported.wps','expected_success':False,'reason':'plain text with WPS suffix must remain unsupported'}, {'filename':'invalid-package.docx','expected_success':False,'reason':'ZIP without document.xml and content types is not DOCX'}]
all_files=[p for p in ROOT.iterdir() if p.is_file() and p.name!='expectations.json']
manifest={'schema_version':1,'purpose':'Synthetic general format and fidelity checks; all source content is fabricated.','fixtures':positive+neg,'baseline_features':{'body_page_break':True,'chinese_and_latin_text':True,'double_ascii_spaces':True,'nbsp':True,'ideographic_space':True,'table_count':1,'table_dimensions':[3,3],'image_count':1,'header_count':1,'footer_count':1,'page_field_count':1,'footnote_count':1,'comment_count':1,'paragraph_styles':['Title','Heading 1','List Number'],'bold_italic_underline':True,'line_break_and_tab':True},'exact_reference_text_file':'baseline-text.json','fidelity_guidance':['Byte copy of valid DOCX, even with a misleading extension, must preserve package bytes.','For conversion formats, verify text, tables, images, header/footer and note/comment parts independently; a success exit code alone is not a preservation guarantee.','NBSP and ordinary spaces are distinct source content; any normalization used for comparison must be explicit.','Do not treat this matrix as exhaustive support for native Kingsoft WPS formats; no genuine Kingsoft proprietary fixture is included.'],'exports':exports,'sha256':{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in all_files}}
(ROOT/'expectations.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'root':str(ROOT),'positive_fixtures':len(positive),'negative_fixtures':len(neg),'exports':exports},ensure_ascii=False,indent=2))
